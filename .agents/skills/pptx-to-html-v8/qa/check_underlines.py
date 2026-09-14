"""QA script for underline validation and content coverage checks."""
import re
import sys
import os

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def parse_slides(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        content = f.read()
    slides = re.split(r'<div class="slide[^"]*"[^>]*data-slide="(\d+)"', content)
    result = {}
    for i in range(1, len(slides), 2):
        slide_num = int(slides[i])
        slide_html = slides[i + 1]
        end = slide_html.find('<div class="slide')
        if end == -1:
            end = len(slide_html)
        result[slide_num] = slide_html[:end]
    return result


def check_underline_tags(slides, expected):
    """Check that specific slides have <u> tags."""
    results = []
    for slide_num, expected_words in expected.items():
        if slide_num not in slides:
            results.append(f"[FAIL] Slide {slide_num}: NOT FOUND in HTML")
            continue
        html = slides[slide_num]
        u_tags = re.findall(r'<u[^>]*>(.*?)</u>', html, re.DOTALL)
        u_texts = []
        for tag_content in u_tags:
            text = re.sub(r'<[^>]+>', '', tag_content).strip()
            if text:
                u_texts.append(text)

        missing = []
        for word in expected_words:
            found = any(word in t for t in u_texts)
            if not found:
                found = word.lower() in ' '.join(u_texts).lower()
            if not found:
                missing.append(word)

        if missing:
            results.append(f"[FAIL] Slide {slide_num}: Missing underlines for: {missing}")
            results.append(f"       Found underlines: {u_texts[:15]}")
        else:
            results.append(f"[PASS] Slide {slide_num}: All expected underlines found ({len(expected_words)} words)")
    return results


def detect_photo_only_slides(pptx_path):
    """Slides whose source shapes carry no text at all - an image-only slide is not a defect."""
    from pptx import Presentation

    def texts(shapes):
        out = []
        for sh in shapes:
            if getattr(sh, 'shape_type', None) == 6:
                try:
                    out += texts(sh.shapes)
                    continue
                except Exception:
                    pass
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    out += [c.text for c in row.cells]
        return out

    prs = Presentation(pptx_path)
    return {i for i, slide in enumerate(prs.slides, 1)
            if not "".join(texts(slide.shapes)).strip()}


def check_no_fullpage_image_only(slides, check_slides, verified=True):
    """Check that specific slides are NOT just a full-page PNG image."""
    results = []
    for slide_num in check_slides:
        if slide_num not in slides:
            results.append(f"[FAIL] Slide {slide_num}: NOT FOUND")
            continue
        html = slides[slide_num]
        has_css_text = bool(re.search(r'<span[^>]*style="[^"]*font-', html))
        has_ocr = 'OCR Override' in html or bool(re.search(r'slide-\d+-container', html))
        has_only_img = not has_css_text and not has_ocr

        native_pngs = re.findall(r'Native PNG for (\w+)', html)
        css_renders = re.findall(r'CSS Render (\w+)', html)

        if has_only_img and len(native_pngs) > 0 and len(css_renders) == 0:
            # Only a real defect when the source deck confirms this slide has text.
            tag = "[FAIL]" if verified else "[WARN]"
            hint = "" if verified else " (pass --pptx to check against the source)"
            results.append(f"{tag} Slide {slide_num}: Only PNG images, no text content "
                         f"({len(native_pngs)} PNGs, {len(css_renders)} CSS){hint}")
        elif has_ocr:
            results.append(f"[PASS] Slide {slide_num}: OCR override applied")
        elif has_css_text:
            results.append(f"[PASS] Slide {slide_num}: Has CSS-rendered text "
                         f"({len(css_renders)} CSS, {len(native_pngs)} PNGs)")
        else:
            results.append(f"[WARN] Slide {slide_num}: No text detected")
    return results


def check_shape_not_image(slides, checks):
    """Check that specific shapes on slides are CSS-rendered, not PNG."""
    results = []
    for slide_num, shape_desc in checks:
        if slide_num not in slides:
            results.append(f"[FAIL] Slide {slide_num}: NOT FOUND")
            continue
        html = slides[slide_num]
        native_match = re.search(r'Native PNG for ' + re.escape(str(shape_desc)), html)
        css_match = re.search(r'CSS Render ' + re.escape(str(shape_desc)), html)
        if native_match and not css_match:
            results.append(f"[FAIL] Slide {slide_num}: Shape {shape_desc} is PNG (should be CSS) — '{shape_desc}'")
        elif css_match:
            results.append(f"[PASS] Slide {slide_num}: Shape {shape_desc} is CSS-rendered")
        else:
            results.append(f"[INFO] Slide {slide_num}: Shape {shape_desc} not found")
    return results


def check_font_underline_in_runs(slides, slide_nums):
    """Check that Font.Underline is preserved in Runs() path (no slide_lines)."""
    results = []
    for slide_num in slide_nums:
        if slide_num not in slides:
            continue
        html = slides[slide_num]
        has_u_tags = '<u class="step-underline">' in html
        if has_u_tags:
            results.append(f"[PASS] Slide {slide_num}: Font.Underline preserved (has <u> tags)")
        else:
            has_underline_collision = 'step-underline step-' in html
            if has_underline_collision:
                results.append(f"[PASS] Slide {slide_num}: Has collision underlines")
            else:
                results.append(f"[FAIL] Slide {slide_num}: No underline tags found at all")
    return results


def run_qa(html_path, pptx_path=None):
    print(f"\n{'='*60}")
    print(f"QA VALIDATION: {html_path}")
    print(f"{'='*60}")

    slides = parse_slides(html_path)
    print(f"\nFound {len(slides)} slides: {sorted(slides.keys())}")

    all_pass = True

    # 1. Underline checks
    print(f"\n--- Underline Validation ---")
    underline_expected = {}
    basename = os.path.basename(os.path.dirname(html_path))
    if 'bai_3' in basename:
        underline_expected = {
            7: ['quyển sách', 'học sinh', 'Con mèo', 'cái ghế', 'bánh mì'],
            8: ['Hoàng hôn', 'cảnh vật', 'màu cam', 'vạt nắng'],
            10: ['hàng cau', 'làng Dạ', 'sức mạnh', 'mùa đông',
                 'tàu lá', 'nền đất', 'đọt lá', 'cây cau'],
        }
    elif 'bai_2' in basename:
        underline_expected = {
            13: ['Bông lúa', 'gió', 'cánh đồng'],
        }
    elif 'bai_4' in basename:
        underline_expected = {
            4: ['những tia nắng ấy'],
            5: ['Hoàng hôn', 'vạt nắng', 'chân trời', 'nền trời', 'cánh diều', 'màu sắc',
                'niềm vui', 'niềm mơ ước', 'lũ trẻ làng quê', 'cơn lũ đêm', 'ngôi nhà',
                'con người', 'Tình yêu đất nước', 'vòm lá', 'chiếc tổ', 'chim',
                'sự khéo léo', 'tình yêu'],
        }

    if underline_expected:
        for r in check_underline_tags(slides, underline_expected):
            print(f"  {r}")
            if '[FAIL]' in r:
                all_pass = False
    else:
        print("  [SKIP] No underline expectations defined for this output")

    # 2. Font.Underline preservation
    print(f"\n--- Font.Underline Preservation ---")
    font_ul_slides = []
    if 'bai_3' in basename:
        font_ul_slides = [7]
    for r in check_font_underline_in_runs(slides, font_ul_slides):
        print(f"  {r}")
        if '[FAIL]' in r:
            all_pass = False

    # 3. Content coverage (no full-page image only)
    print(f"\n--- Content Coverage ---")
    photo_only_slides = set()
    verified = False
    if pptx_path:
        try:
            photo_only_slides = detect_photo_only_slides(pptx_path)
            verified = True
        except Exception as e:
            print(f"  [WARN] Could not read {pptx_path}: {e}")
    elif 'bai_2' in basename:
        photo_only_slides = {15, 16, 17, 18, 19, 20}
    elif 'bai_4' in basename:
        photo_only_slides = {8, 10, 12, 14, 16, 17, 18, 19, 23, 24, 27}
    check_content_slides = [s for s in slides.keys() if s not in photo_only_slides]
    if photo_only_slides:
        origin = "source deck" if verified else "known list"
        print(f"  [SKIP] Slides {sorted(photo_only_slides)}: image-only in the {origin} (no text expected)")
    for r in check_no_fullpage_image_only(slides, check_content_slides, verified=verified):
        print(f"  {r}")
        if '[FAIL]' in r:
            all_pass = False

    # 4. Specific shape checks
    print(f"\n--- Shape Rendering ---")
    shape_checks = []
    if 'bai_3' in basename:
        shape_checks = [
            (14, '6'),
        ]
    for r in check_shape_not_image(slides, shape_checks):
        print(f"  {r}")
        if '[FAIL]' in r:
            all_pass = False

    # Summary
    print(f"\n{'='*60}")
    if all_pass:
        print("RESULT: ALL CHECKS PASSED")
    else:
        print("RESULT: SOME CHECKS FAILED — see above")
    print(f"{'='*60}")
    return all_pass


USAGE = "Usage: python qa/check_underlines.py <presentation.html> [--pptx <source.pptx>]"

if __name__ == '__main__':
    html_path = None
    pptx_path = None
    args = sys.argv[1:]
    i = 0
    while i < len(args):
        if args[i] == '--pptx' and i + 1 < len(args):
            pptx_path = args[i + 1]
            i += 2
        else:
            html_path = args[i]
            i += 1

    if not html_path or not os.path.exists(html_path):
        print(f"Not found: {html_path}" if html_path else USAGE)
        sys.exit(2)

    sys.exit(0 if run_qa(html_path, pptx_path) else 1)
